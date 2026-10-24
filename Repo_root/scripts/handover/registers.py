#!/usr/bin/env python3
"""Write the Folio Registers workbook for the handover pack.

    python3 scripts/handover/registers.py <input.json> <output.xlsx>

The input is written by scripts/build-handover.mjs: a title block and a list of sheets, each with its columns
(name, width) and rows. This file only lays them out; every value comes from the registers.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GREEN = '05583B'
TINT = 'E6EEEB'
RULE = Side(style='thin', color='CDDDD7')

data = json.load(open(sys.argv[1], encoding='utf-8'))
wb = Workbook()

# Cover sheet: the title block and what each sheet holds.
cover = wb.active
cover.title = 'About'
cover['A1'] = data['title']
cover['A1'].font = Font(name='Inter', size=16, bold=True, color=GREEN)
row = 3
for label, value in data['control']:
    cover.cell(row=row, column=1, value=label).font = Font(name='Inter', bold=True)
    cover.cell(row=row, column=2, value=value).font = Font(name='Inter')
    row += 1
row += 1
cover.cell(row=row, column=1, value='Sheet').font = Font(name='Inter', bold=True, color='FFFFFF')
cover.cell(row=row, column=2, value='What it holds').font = Font(name='Inter', bold=True, color='FFFFFF')
for c in (1, 2):
    cover.cell(row=row, column=c).fill = PatternFill('solid', fgColor=GREEN)
for sheet in data['sheets']:
    row += 1
    cover.cell(row=row, column=1, value=sheet['name']).font = Font(name='Inter')
    cover.cell(row=row, column=2, value=sheet['about']).font = Font(name='Inter')
    cover.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical='top')
cover.column_dimensions['A'].width = 28
cover.column_dimensions['B'].width = 110

for sheet in data['sheets']:
    ws = wb.create_sheet(sheet['name'][:31])
    for i, (name, width) in enumerate(sheet['columns'], start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.font = Font(name='Inter', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=GREEN)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    for r, values in enumerate(sheet['rows'], start=2):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name='Inter', size=10)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.border = Border(bottom=RULE)
            if r % 2 == 0:
                cell.fill = PatternFill('solid', fgColor=TINT)
    ws.freeze_panes = 'A2'
    if sheet['rows']:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(sheet['columns']))}{len(sheet['rows']) + 1}"
    ws.sheet_view.zoomScale = 90

wb.save(sys.argv[2])
print(f"{sys.argv[2]}: {len(data['sheets'])} sheets, {sum(len(s['rows']) for s in data['sheets'])} rows")
