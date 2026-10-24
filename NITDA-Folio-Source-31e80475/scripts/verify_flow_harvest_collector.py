#!/usr/bin/env python3
"""Prove the collector workbook is actually writable by a Power Automate flow.

The failure this guards against is silent and expensive. "Add a row into a table" maps values
to columns BY NAME, against a table the connector must be able to see. A sheet that looks
correct in Excel but carries an AutoFilter instead of a ListObject is invisible to the
connector; a column renamed by hand breaks one mapping and leaves the rest working; two columns
sharing a name make the mapping ambiguous. None of that shows up until a harvest has run and
the data is wrong or missing.

So this opens the delivered file — not the spec, the file — and checks what a flow author would
otherwise discover the hard way:

  · every table exists as a real ListObject, addressable by the name the flow will use
  · its columns are exactly the specified ones, in order, spelled identically
  · no two columns in a table share a name
  · the table range starts at the header row and includes the placeholder row Excel needs
  · the placeholder row is genuinely empty, so no harvest starts with a phantom record
  · the data dictionary describes every column that exists and no column that does not

Usage: python3 scripts/verify_flow_harvest_collector.py [workbook.xlsx]
Exit:  0 = a flow can write to it, 1 = otherwise
"""

import sys

from openpyxl import load_workbook

sys.path.insert(0, "scripts")
from build_flow_harvest_collector import (  # noqa: E402
    SPEC, DICT_COLUMNS, HEADER_ROW, PLACEHOLDER_ROW, VERSION,
)

EXCEL_NAME_LIMIT = 255


def main(path):
    wb = load_workbook(path)
    failures = []
    checked = 0

    def check(cond, message):
        nonlocal checked
        checked += 1
        if not cond:
            failures.append(message)

    tables_seen = {}
    for ws in wb.worksheets:
        for name, ref in ws.tables.items():
            tables_seen[name] = (ws, ref)

    expected = {s["table"]: s for s in SPEC}
    expected_names = set(expected) | {"tblDictionary"}
    check(set(tables_seen) == expected_names,
          f"tables in the file {sorted(tables_seen)} do not match the spec {sorted(expected_names)}")

    for spec in SPEC:
        name = spec["table"]
        if name not in tables_seen:
            failures.append(f"{name}: not a named table in the workbook — the connector cannot see it")
            checked += 1
            continue
        ws, ref = tables_seen[name]
        table = ws.tables[name]
        cols = [c["name"] for c in spec["columns"]]

        header = [ws.cell(row=HEADER_ROW, column=i).value for i in range(1, len(cols) + 1)]
        check(header == cols,
              f"{name}: header row is {header!r}, spec says {cols!r}")

        listed = [c.name for c in table.tableColumns]
        check(listed == cols,
              f"{name}: the table's own column list is {listed!r}, spec says {cols!r}")

        check(len(set(cols)) == len(cols),
              f"{name}: duplicate column name — 'Add a row into a table' maps by name and cannot resolve it")
        check(all(c and c.strip() == c and len(c) <= EXCEL_NAME_LIMIT for c in cols),
              f"{name}: a column name is empty, padded, or too long for Excel")

        check(ref.startswith(f"A{HEADER_ROW}:"),
              f"{name}: range {ref} does not start at the header row A{HEADER_ROW}")
        check(ref.endswith(str(PLACEHOLDER_ROW)),
              f"{name}: range {ref} does not end at the placeholder row {PLACEHOLDER_ROW} — "
              f"a header-only table makes Excel repair the file, and extra rows are phantom records")

        blank = [ws.cell(row=PLACEHOLDER_ROW, column=i).value for i in range(1, len(cols) + 1)]
        check(all(v is None for v in blank),
              f"{name}: the placeholder row carries values {[v for v in blank if v is not None]!r} — "
              f"a harvest would start with a record nobody wrote")

    # ── the dictionary must describe the tables, exactly ──────────────────────────────────
    if "tblDictionary" in tables_seen:
        ws, _ = tables_seen["tblDictionary"]
        header = [ws.cell(row=HEADER_ROW, column=i).value for i in range(1, len(DICT_COLUMNS) + 1)]
        check(header == DICT_COLUMNS, f"tblDictionary: header is {header!r}")

        rows = []
        r = PLACEHOLDER_ROW + 1
        while ws.cell(row=r, column=1).value is not None:
            rows.append(tuple(ws.cell(row=r, column=i).value for i in range(1, len(DICT_COLUMNS) + 1)))
            r += 1
        want = [(s["table"], s["sheet"], c["name"], c["type"], c["filled_by"], c["source"], c["meaning"])
                for s in SPEC for c in s["columns"]]
        check(rows == want,
              f"tblDictionary: {len(rows)} entries, spec has {len(want)}"
              + ("" if len(rows) == len(want) else "")
              + next((f"; first difference at {i}: {g!r} vs {w!r}"
                      for i, (g, w) in enumerate(zip(rows, want)) if g != w), ""))
        described = {(t, c) for t, _s, c, *_ in rows}
        actual = {(s["table"], c["name"]) for s in SPEC for c in s["columns"]}
        check(described == actual, "tblDictionary describes columns that do not exist, or misses some")

    print(f"  tables required           {len(expected_names)}")
    print(f"  tables found as ListObjects {len(tables_seen)}")
    print(f"  columns under contract    {sum(len(s['columns']) for s in SPEC)}")
    print(f"  assertions run            {checked}")
    print(f"  workbook contract version {VERSION}")

    if failures:
        print(f"\n❌ {len(failures)} problem(s) — a flow cannot reliably write to this workbook:")
        for f in failures:
            print(f"     {f}")
        return 1

    print("\n✅ every table is a real named table with exactly the columns the mapping expects,")
    print("   each one empty apart from the placeholder row Excel requires. A flow can write to it.")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else "dist/DGO_Flow_Harvest_Collector.xlsx"))
